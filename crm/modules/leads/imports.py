"""Lead import wizard (Part 4.8): upload -> map columns -> validate (dry run) -> import -> result, with 24 h undo."""
from __future__ import annotations

import csv
import io
import re

from django.contrib import messages
from django.core.exceptions import ValidationError
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST

from crm.core import audit, crypto, india, permissions
from crm.core.views import staff_required
from crm.models import ImportBatch, Lead, LeadActivity, Service

from . import services as svc
from .forms import ImportOptionsForm, ImportUploadForm

MAX_ROWS = 10_000
FIELDS = [
    ('name', 'Name', True), ('phone', 'Phone', False), ('email', 'Email', False), ('company', 'Company', False),
    ('city', 'City', False), ('notes', 'Notes', False), ('service', 'Service interested', False),
    ('priority', 'Priority (hot/warm/cold)', False), ('value', 'Expected value', False),
]
GUESS = {
    'name': ['name', 'full name', 'contact name', 'lead name', 'customer name'],
    'phone': ['phone', 'mobile', 'phone number', 'mobile number', 'contact number', 'whatsapp'],
    'email': ['email', 'email address', 'e-mail'],
    'company': ['company', 'company name', 'business', 'organisation', 'organization'],
    'city': ['city', 'location', 'place'],
    'notes': ['notes', 'message', 'requirement', 'comments', 'remarks'],
    'service': ['service', 'interest', 'service interested', 'product'],
    'priority': ['priority'],
    'value': ['value', 'budget', 'deal value', 'amount'],
}
TEMPLATE_HEADERS = ['Name', 'Phone', 'Email', 'Company', 'City', 'Notes', 'Service', 'Priority', 'Value']


def read_rows(batch: ImportBatch) -> tuple[list[str], list[list[str]]]:
    batch.file.open('rb')
    data = batch.file.read()
    batch.file.close()
    if batch.file_name.lower().endswith('.xlsx'):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = [['' if v is None else str(v).strip() for v in r] for r in ws.iter_rows(values_only=True)]
    else:
        text = data.decode('utf-8-sig', errors='replace')
        rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text))]
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        raise ValidationError('The file is empty.')
    headers, body = rows[0], rows[1:]
    if len(body) > MAX_ROWS:
        raise ValidationError(f'The file has {len(body)} rows. The limit is {MAX_ROWS:,}.')
    return headers, body


def guess_mapping(headers: list[str]) -> dict:
    out = {}
    norm = [re.sub(r'[^a-z ]', '', h.lower()).strip() for h in headers]
    for field, options in GUESS.items():
        for i, h in enumerate(norm):
            if h in options and str(i) not in out.values():
                out[field] = str(i)
                break
    return out


def validate_rows(batch: ImportBatch) -> list[dict]:
    headers, body = read_rows(batch)
    m = {k: int(v) for k, v in batch.mapping.items() if v != ''}
    results = []
    seen_phone, seen_email = set(), set()
    services = {s.name.lower(): s.pk for s in Service.objects.filter(is_active=True)}
    for n, row in enumerate(body, start=2):
        get = lambda f: (row[m[f]].strip() if f in m and m[f] < len(row) else '')  # noqa: E731
        res = {'row': n, 'status': 'valid', 'errors': [], 'data': {}}
        name = get('name')
        phone_raw, email = get('phone'), get('email').lower()
        phone = ''
        if phone_raw:
            try:
                phone = india.normalize_phone(phone_raw)
            except ValueError as exc:
                res['errors'].append(f'Phone: {exc}')
        if email and not re.match(r'^[^@\s]+@[^@\s]+\.[^@\s]+$', email):
            res['errors'].append('Email is not valid')
            email = ''
        if len(name) < 2:
            res['errors'].append('Name must be at least 2 characters')
        if not phone and not email:
            res['errors'].append('Needs a phone number or email')
        priority = get('priority').lower()
        if priority and priority not in ('hot', 'warm', 'cold'):
            res['errors'].append('Priority must be hot, warm or cold')
        value = get('value').replace(',', '').replace('₹', '')
        if value:
            try:
                float(value)
            except ValueError:
                res['errors'].append('Value is not a number')
                value = ''
        svc_name = get('service').lower()
        res['data'] = {'name': name[:100], 'phone': phone, 'email': email, 'company': get('company')[:150],
                       'city': get('city')[:80], 'notes': get('notes')[:2000], 'priority': priority or 'warm',
                       'value': value, 'service_id': services.get(svc_name) if svc_name else None}
        if res['errors']:
            res['status'] = 'invalid'
        else:
            key_p, key_e = phone or None, email or None
            if (key_p and key_p in seen_phone) or (key_e and key_e in seen_email):
                res['status'] = 'duplicate'
                res['errors'].append('Repeated within this file')
            elif svc.find_duplicates(phone, email):
                res['status'] = 'duplicate'
                res['errors'].append('Matches an existing lead')
            if key_p:
                seen_phone.add(key_p)
            if key_e:
                seen_email.add(key_e)
        results.append(res)
    return results


@staff_required
def import_upload(request):
    permissions.require(request.user, 'leads', 'import', request)
    form = ImportUploadForm(request.POST or None, request.FILES or None)
    if request.method == 'POST' and form.is_valid():
        f = form.cleaned_data['file']
        digest = crypto.sha256_file(f)
        batch = ImportBatch.objects.create(module='leads', file_name=f.name[:255], file_hash=digest, file=f,
                                           created_by=request.user, status='uploaded')
        try:
            headers, body = read_rows(batch)
        except ValidationError as exc:
            batch.status = 'failed'
            batch.save(update_fields=['status'])
            form.add_error('file', ' '.join(exc.messages))
        else:
            batch.headers = headers
            batch.total_rows = len(body)
            batch.mapping = guess_mapping(headers)
            batch.save(update_fields=['headers', 'total_rows', 'mapping'])
            audit.log('import.batch_created', batch, changes={'file_hash': digest, 'rows': len(body)})
            return redirect('crm:lead_import_map', pk=batch.pk)
    recent = ImportBatch.objects.filter(module='leads').select_related('created_by')[:10]
    return render(request, 'crm/leads/import_upload.html', {'form': form, 'recent': recent, 'step': 1})


def _batch(request, pk):
    permissions.require(request.user, 'leads', 'import', request)
    batch = get_object_or_404(ImportBatch, pk=pk, module='leads')
    if batch.created_by_id != request.user.pk and not permissions.has_perm(request.user, 'leads', 'approve'):
        permissions.denied(request.user, 'leads', 'import', request, detail=f'batch {pk}')
    return batch


@staff_required
def import_map(request, pk):
    batch = _batch(request, pk)
    if batch.status in ('imported', 'undone'):
        return redirect('crm:lead_import_result', pk=batch.pk)
    error = ''
    if request.method == 'POST':
        mapping = {f: request.POST.get(f'map_{f}', '') for f, _, _ in FIELDS}
        if mapping.get('name', '') == '':
            error = 'Map the Name column.'
        elif mapping.get('phone', '') == '' and mapping.get('email', '') == '':
            error = 'Map a Phone or Email column.'
        else:
            batch.mapping = mapping
            batch.status = 'mapped'
            batch.save(update_fields=['mapping', 'status'])
            return redirect('crm:lead_import_validate', pk=batch.pk)
    headers, body = read_rows(batch)
    sample = body[:3]
    return render(request, 'crm/leads/import_map.html', {'batch': batch, 'fields': FIELDS, 'headers': list(enumerate(headers)),
                                                         'sample': sample, 'error': error, 'step': 2})


@staff_required
def import_validate(request, pk):
    batch = _batch(request, pk)
    if batch.status in ('imported', 'undone'):
        return redirect('crm:lead_import_result', pk=batch.pk)
    results = validate_rows(batch)
    batch.valid_rows = sum(1 for r in results if r['status'] == 'valid')
    batch.duplicate_rows = sum(1 for r in results if r['status'] == 'duplicate')
    batch.invalid_rows = sum(1 for r in results if r['status'] == 'invalid')
    batch.row_results = results
    batch.status = 'validated'
    batch.save(update_fields=['valid_rows', 'duplicate_rows', 'invalid_rows', 'row_results', 'status'])
    if request.GET.get('download') == 'errors':
        out = io.StringIO()
        w = csv.writer(out)
        w.writerow(['Row', 'Status', 'Problems'])
        for r in results:
            if r['status'] != 'valid':
                w.writerow([r['row'], r['status'], '; '.join(r['errors'])])
        resp = HttpResponse('﻿' + out.getvalue(), content_type='text/csv; charset=utf-8')
        resp['Content-Disposition'] = f'attachment; filename="import-{batch.pk}-errors.csv"'
        return resp
    form = ImportOptionsForm(user=request.user)
    return render(request, 'crm/leads/import_validate.html', {'batch': batch, 'form': form, 'problems': [r for r in results if r['status'] != 'valid'][:200], 'step': 3})


@staff_required
@require_POST
def import_run(request, pk):
    batch = _batch(request, pk)
    if batch.status == 'imported':
        messages.info(request, 'This batch was already imported; nothing was repeated.')
        return redirect('crm:lead_import_result', pk=batch.pk)
    if batch.status != 'validated':
        return redirect('crm:lead_import_validate', pk=batch.pk)
    form = ImportOptionsForm(request.POST, user=request.user)
    if not form.is_valid():
        messages.error(request, 'Choose how to handle duplicates and the source.')
        return redirect('crm:lead_import_validate', pk=batch.pk)
    opts = form.cleaned_data
    results = validate_rows(batch)   # re-validate: data may have changed since the dry run
    created = interaction = skipped = 0
    with transaction.atomic():
        batch = ImportBatch.objects.select_for_update().get(pk=batch.pk)
        if batch.status == 'imported':
            return redirect('crm:lead_import_result', pk=batch.pk)
        for r in results:
            d = r['data']
            if r['status'] == 'invalid':
                skipped += 1
                r['result'] = 'skipped (invalid)'
                continue
            try:
                with transaction.atomic():
                    if r['status'] == 'duplicate':
                        existing = svc.find_duplicates(d['phone'], d['email'])
                        if opts['duplicates'] == 'skip' or (not existing and opts['duplicates'] == 'interaction'):
                            skipped += 1
                            r['result'] = 'skipped (duplicate)'
                            continue
                        if opts['duplicates'] == 'interaction' and existing:
                            svc.attach_enquiry(existing[0], f'Imported again (batch #{batch.pk}): {d["notes"] or "no notes"}',
                                               source=opts['source'], actor=request.user)
                            interaction += 1
                            r['result'] = f'added to lead #{existing[0].pk}'
                            continue
                    lead = svc.create_lead(svc.LeadInput(
                        name=d['name'], phone=d['phone'], email=d['email'], company=d['company'], city=d['city'],
                        notes=d['notes'], source=opts['source'], campaign=opts.get('campaign'), owner=opts.get('owner'),
                        priority=d['priority'], value=d['value'] or None), request.user,
                        services=[d['service_id']] if d.get('service_id') else [],
                        rule='import', import_batch=batch,
                        duplicate_reason=f'Imported anyway (batch #{batch.pk})' if r['status'] == 'duplicate' else '')
                    created += 1
                    r['result'] = f'created lead #{lead.pk}'
                    r['lead_id'] = lead.pk
            except ValidationError as exc:
                skipped += 1
                r['result'] = 'failed: ' + ' '.join(exc.messages)
        batch.row_results = results
        batch.created_count, batch.interaction_count, batch.skipped_count = created, interaction, skipped
        batch.status = 'imported'
        batch.imported_at = timezone.now()
        batch.options = {'duplicates': opts['duplicates'], 'source': opts['source'].pk,
                         'owner': opts['owner'].pk if opts.get('owner') else None}
        batch.save()
        audit.log('import.completed', batch, changes={'created': created, 'interactions': interaction, 'skipped': skipped})
    messages.success(request, f'Import finished: {created} created, {interaction} added to existing leads, {skipped} skipped.')
    return redirect('crm:lead_import_result', pk=batch.pk)


@staff_required
def import_result(request, pk):
    batch = _batch(request, pk)
    return render(request, 'crm/leads/import_result.html', {'batch': batch, 'rows': batch.row_results[:500], 'step': 5})


@staff_required
@require_POST
def import_undo(request, pk):
    """Undo within 24 h for leads with no human activity; the batch itself is never deleted."""
    batch = _batch(request, pk)
    if not batch.can_undo:
        messages.error(request, 'Undo is only available for 24 hours after import.')
        return redirect('crm:lead_import_result', pk=batch.pk)
    archived = kept = 0
    with transaction.atomic():
        for lead in Lead.objects.select_for_update().filter(import_batch=batch, archived_at__isnull=True):
            human = LeadActivity.objects.filter(lead=lead, actor__isnull=False).exclude(kind__in=['assignment', 'system', 'inbound']).exists()
            if human or lead.quotations.exists():
                kept += 1
                continue
            lead.archived_at = timezone.now()
            lead.archived_by = request.user
            lead.archive_reason = f'Import #{batch.pk} undone'
            lead.save(update_fields=['archived_at', 'archived_by', 'archive_reason'])
            lead.follow_ups.filter(status='pending').update(status='cancelled', notes='Import undone')
            archived += 1
        batch.undone_at = timezone.now()
        batch.status = 'undone'
        batch.save(update_fields=['undone_at', 'status'])
        audit.log('import.undone', batch, changes={'archived': archived, 'kept_with_activity': kept})
    messages.success(request, f'Undo complete: {archived} leads archived, {kept} kept because work had started on them.')
    return redirect('crm:lead_import_result', pk=batch.pk)


@staff_required
def import_template(request):
    permissions.require(request.user, 'leads', 'import', request)
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(TEMPLATE_HEADERS)
    w.writerow(['Anita K.', '9876543210', 'anita@example.com', 'Anita Textiles', 'Kochi', 'Wants an online store',
                'Online Store', 'warm', '35000'])
    resp = HttpResponse('﻿' + out.getvalue(), content_type='text/csv; charset=utf-8')
    resp['Content-Disposition'] = 'attachment; filename="lead-import-template.csv"'
    return resp
