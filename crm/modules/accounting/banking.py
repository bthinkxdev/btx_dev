"""Bank statement import, matching and reconciliation (Part 8.5)."""
from __future__ import annotations

import csv
import hashlib
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.db.models import Sum

from crm.core import audit, crypto, permissions, posting
from crm.core.money import ZERO, D, q2
from crm.models import (Account, BankAccount, BankReconciliation, BankStatementImport, BankTransaction, JournalLine, Payment,
                        ReconciliationMatch, Settlement)


class BankError(ValidationError):
    pass


DATE_FORMATS = ('%d-%m-%Y', '%d/%m/%Y', '%Y-%m-%d', '%d-%b-%Y', '%d %b %Y', '%d-%b-%y', '%d/%m/%y', '%d.%m.%Y', '%m/%d/%Y')
HEADER_ALIASES = {
    'date': ('date', 'txn date', 'transaction date', 'value date', 'posting date', 'tran date'),
    'description': ('description', 'narration', 'particulars', 'details', 'remarks', 'transaction remarks'),
    'reference': ('reference', 'ref', 'ref no', 'ref no.', 'chq / ref no.', 'cheque no', 'chq no', 'utr', 'reference no'),
    'debit': ('debit', 'withdrawal', 'withdrawal amt', 'withdrawal amount', 'dr', 'debit amount', 'withdrawals'),
    'credit': ('credit', 'deposit', 'deposit amt', 'deposit amount', 'cr', 'credit amount', 'deposits'),
    'amount': ('amount', 'transaction amount'),
    'balance': ('balance', 'closing balance', 'running balance', 'available balance'),
}


def parse_date(raw) -> date:
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    s = str(raw).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(f'Unrecognised date "{s}"')


def parse_amount(raw) -> Decimal:
    if raw is None:
        return ZERO
    if isinstance(raw, (int, float, Decimal)):
        return q2(D(str(raw)))
    s = str(raw).strip().replace(',', '').replace('₹', '').replace('INR', '').strip()
    if not s or s in ('-', '—'):
        return ZERO
    neg = s.endswith('Dr') or s.startswith('(') or s.startswith('-')
    s = re.sub(r'[^\d.]', '', s)
    if not s:
        return ZERO
    try:
        v = q2(Decimal(s))
    except InvalidOperation as exc:
        raise ValueError(f'Unrecognised amount "{raw}"') from exc
    return -v if neg else v


def _read_rows(upload) -> list[list]:
    name = upload.name.lower()
    if name.endswith('.xlsx'):
        from openpyxl import load_workbook
        wb = load_workbook(upload, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        return [[c for c in row] for row in ws.iter_rows(values_only=True)]
    text = upload.read().decode('utf-8-sig', errors='replace')
    return list(csv.reader(io.StringIO(text)))


def _map_header(rows):
    """Find the header row (banks put a preamble above it) and map columns."""
    for idx, row in enumerate(rows[:30]):
        cells = [str(c or '').strip().lower() for c in row]
        mapping = {}
        for key, aliases in HEADER_ALIASES.items():
            for i, c in enumerate(cells):
                if c in aliases and key not in mapping:
                    mapping[key] = i
        if 'date' in mapping and 'description' in mapping and ('amount' in mapping or 'debit' in mapping or 'credit' in mapping):
            return idx, mapping
    raise BankError('Could not find the header row. It needs Date, Description and Amount (or Debit and Credit) columns.')


def line_hash(d: date, amount: Decimal, description: str, balance) -> str:
    key = f'{d.isoformat()}|{amount}|{" ".join(description.lower().split())}|{balance if balance is not None else ""}'
    return hashlib.sha256(key.encode()).hexdigest()


def import_statement(bank: BankAccount, upload, actor) -> BankStatementImport:
    from crm.core.forms import validate_upload
    validate_upload(upload, ('.csv', '.xlsx'))
    digest = crypto.sha256_file(upload)
    rows = _read_rows(upload)
    if not rows:
        raise BankError('The file is empty.')
    head, m = _map_header(rows)
    errors, created, dups = [], 0, 0
    with transaction.atomic():
        imp = BankStatementImport.objects.create(bank_account=bank, file_name=upload.name[:255], file_hash=digest, imported_by=actor)
        for n, row in enumerate(rows[head + 1:], start=head + 2):
            if not row or not any(str(c or '').strip() for c in row):
                continue
            try:
                raw_date = row[m['date']] if m['date'] < len(row) else ''
                if not str(raw_date or '').strip():
                    continue
                d = parse_date(raw_date)
                desc = str(row[m['description']] or '').strip()[:500]
                if 'amount' in m:
                    amount = parse_amount(row[m['amount']])
                else:
                    dr = parse_amount(row[m['debit']]) if 'debit' in m and m['debit'] < len(row) else ZERO
                    cr = parse_amount(row[m['credit']]) if 'credit' in m and m['credit'] < len(row) else ZERO
                    amount = abs(cr) - abs(dr)
                if amount == 0:
                    continue
                bal = parse_amount(row[m['balance']]) if 'balance' in m and m['balance'] < len(row) and row[m['balance']] not in (None, '') else None
                ref = str(row[m['reference']] or '').strip()[:120] if 'reference' in m and m['reference'] < len(row) else ''
            except (ValueError, IndexError) as exc:
                errors.append({'row': n, 'error': str(exc)})
                continue
            h = line_hash(d, amount, desc, bal)
            try:
                with transaction.atomic():
                    BankTransaction.objects.create(bank_account=bank, statement_import=imp, txn_date=d, description=desc or '(no narration)',
                                                   reference=ref, amount=amount, running_balance=bal, line_hash=h)
                created += 1
            except IntegrityError:
                dups += 1
        imp.line_count, imp.duplicate_count, imp.error_count, imp.errors = created, dups, len(errors), errors[:200]
        imp.save(update_fields=['line_count', 'duplicate_count', 'error_count', 'errors'])
        audit.log('bank.statement_imported', imp, actor=actor, changes={'account': str(bank), 'lines': created, 'duplicates': dups,
                                                                       'errors': len(errors), 'sha256': digest})
    auto_match(bank, actor)
    return imp


# ---------------------------------------------------------------- matching

def open_book_lines(bank: BankAccount, until: date | None = None):
    """Posted ledger lines on the bank's account that are not matched to a bank line yet."""
    qs = JournalLine.objects.filter(account=bank.ledger_account, posted=True, bank_match__isnull=True).select_related('entry')
    if until:
        qs = qs.filter(entry_date__lte=until)
    return qs


def _signed(ln: JournalLine) -> Decimal:
    return ln.debit - ln.credit


def _references_for(ln: JournalLine) -> list[str]:
    e = ln.entry
    refs = []
    if e.source_type == 'payment':
        p = Payment.objects.filter(pk=e.source_id).first()
        if p:
            refs += [p.reference_number, p.provider_payment_id, p.receipt_number]
    elif e.source_type == 'settlement':
        s = Settlement.objects.filter(pk=e.source_id).first()
        if s:
            refs.append(s.settlement_id)
    elif e.source_type == 'expensepayment':
        from crm.models import ExpensePayment
        ep = ExpensePayment.objects.filter(pk=e.source_id).first()
        if ep:
            refs.append(ep.reference)
    elif e.source_type == 'tdschallan':
        from crm.models import TdsChallan
        c = TdsChallan.objects.filter(pk=e.source_id).first()
        if c:
            refs.append(c.cin)
    return [r.lower() for r in refs if r and len(r) >= 4]


def _match(txn: BankTransaction, lines: list[JournalLine], actor, rule: str):
    for ln in lines:
        ReconciliationMatch.objects.create(bank_transaction=txn, journal_line=ln, amount=_signed(ln), matched_by=actor, rule=rule)
    txn.status = 'matched'
    txn.suggestion = {}
    txn.save(update_fields=['status', 'suggestion'])


def auto_match(bank: BankAccount, actor=None) -> dict:
    """Rules in order: reference + amount, settlement id, single amount within 3 days (suggested)."""
    matched = suggested = 0
    with transaction.atomic():
        txns = list(BankTransaction.objects.select_for_update().filter(bank_account=bank, status__in=['unmatched', 'suggested']))
        pool = list(open_book_lines(bank))
        used = set()
        for t in txns:
            text = f'{t.reference} {t.description}'.lower()
            cands = [ln for ln in pool if ln.pk not in used and _signed(ln) == t.amount]
            hit = next((ln for ln in cands if any(r in text for r in _references_for(ln))), None)
            if hit:
                _match(t, [hit], actor, 'reference')
                used.add(hit.pk)
                matched += 1
                continue
            near = [ln for ln in cands if abs((ln.entry_date - t.txn_date).days) <= 3]
            if len(near) == 1:
                t.status = 'suggested'
                t.suggestion = {'lines': [near[0].pk], 'rule': 'amount_date',
                                'label': f'{near[0].entry.number} · {near[0].entry.narration[:80]}'}
                t.save(update_fields=['status', 'suggestion'])
                suggested += 1
    return {'matched': matched, 'suggested': suggested}


def confirm_suggestion(txn: BankTransaction, actor):
    ids = (txn.suggestion or {}).get('lines') or []
    if txn.status != 'suggested' or not ids:
        raise BankError('There is no suggestion to confirm.')
    return manual_match(txn, ids, actor, rule='suggested')


def manual_match(txn: BankTransaction, line_ids: list[int], actor, *, rule: str = 'manual', adjustment_account: Account | None = None,
                 reason: str = ''):
    """Match one bank line to one or more book items; a difference is posted to an adjustment account with a reason."""
    with transaction.atomic():
        txn = BankTransaction.objects.select_for_update().get(pk=txn.pk)
        if txn.status in ('matched', 'reconciled', 'ignored'):
            raise BankError('This bank line is already matched or closed.')
        lines = list(JournalLine.objects.select_for_update().filter(pk__in=line_ids, account=txn.bank_account.ledger_account,
                                                                    posted=True, bank_match__isnull=True))
        if len(lines) != len(set(line_ids)):
            raise BankError('Some selected book items are already matched or belong to another account.')
        total = sum((_signed(ln) for ln in lines), ZERO)
        diff = txn.amount - total
        if diff:
            if adjustment_account is None or len(reason.strip()) < 5:
                raise BankError(f'The selection differs from the bank line by ₹{diff:,.2f}. Choose an adjustment account and give a reason.')
            je = posting.build_journal(entry_date=txn.txn_date, narration=f'Bank difference: {reason}'[:500], source_type='banktransaction',
                                       source_id=str(txn.pk), event_key=f'adjustment:{len(line_ids)}', user=actor,
                                       lines=[posting.Line(txn.bank_account.ledger_account, debit=diff if diff > 0 else ZERO,
                                                           credit=-diff if diff < 0 else ZERO, description=txn.description[:300]),
                                              posting.Line(adjustment_account, debit=-diff if diff < 0 else ZERO,
                                                           credit=diff if diff > 0 else ZERO, description=reason[:300])])
            lines.append(je.lines.get(account=txn.bank_account.ledger_account))
        _match(txn, lines, actor, rule)
        audit.log('bank.matched', txn, actor=actor, changes={'lines': [ln.pk for ln in lines], 'rule': rule, 'difference': str(diff)},
                  reason=reason)
    return txn


def create_from_line(txn: BankTransaction, actor, *, account: Account, narration: str = ''):
    """Unmatched bank line -> bank charge, interest, drawings, capital, transfer or other entry, matched at once."""
    if txn.status not in ('unmatched', 'suggested'):
        raise BankError('Only an unmatched bank line can create an entry.')
    if account.pk == txn.bank_account.ledger_account_id:
        raise BankError('Choose the other side of the entry, not the bank account itself.')
    with transaction.atomic():
        row = posting.enqueue('bank_line_entry', txn, 'entry', {'account_id': account.pk, 'narration': narration or txn.description[:300]})
        posting.process(row.pk)
        row.refresh_from_db()
        if row.status != 'posted':
            raise BankError(f'Could not post: {row.last_error}')
        ln = row.journal.lines.get(account=txn.bank_account.ledger_account)
        txn = BankTransaction.objects.select_for_update().get(pk=txn.pk)
        _match(txn, [ln], actor, 'created')
        audit.log('bank.entry_created', txn, actor=actor, changes={'account': str(account), 'journal': row.journal.number})
    return row.journal


def ignore_line(txn: BankTransaction, actor, reason: str):
    if len(reason.strip()) < 10:
        raise BankError('Give a reason of at least 10 characters.')
    if txn.status not in ('unmatched', 'suggested'):
        raise BankError('Only an unmatched line can be ignored.')
    BankTransaction.objects.filter(pk=txn.pk).update(status='ignored', ignore_reason=reason[:300], suggestion={})
    audit.log('bank.ignored', txn, actor=actor, reason=reason)


def unmatch(txn: BankTransaction, actor, reason: str):
    if not permissions.has_perm(actor, 'accounting', 'approve'):
        permissions.denied(actor, 'accounting', 'approve', detail='unmatch bank line')
    if len(reason.strip()) < 10:
        raise BankError('Give a reason of at least 10 characters.')
    with transaction.atomic():
        txn = BankTransaction.objects.select_for_update().get(pk=txn.pk)
        if txn.status == 'reconciled':
            raise BankError('Un-reconcile the statement first.')
        if txn.status == 'ignored':
            BankTransaction.objects.filter(pk=txn.pk).update(status='unmatched', ignore_reason='')
        else:
            txn.matches.all().delete()
            BankTransaction.objects.filter(pk=txn.pk).update(status='unmatched')
        audit.log('bank.unmatched', txn, actor=actor, reason=reason)


# ---------------------------------------------------------------- reconciliation

def reconciliation_position(bank: BankAccount, as_of: date, statement_balance: Decimal) -> dict:
    book = posting.balance(bank.ledger_account, as_of=as_of)
    outstanding_lines = list(open_book_lines(bank, as_of))
    outstanding = sum((_signed(ln) for ln in outstanding_lines), ZERO)
    unmatched_bank = BankTransaction.objects.filter(bank_account=bank, txn_date__lte=as_of, status__in=['unmatched', 'suggested'])
    unmatched_total = unmatched_bank.aggregate(s=Sum('amount'))['s'] or ZERO
    # statement = book - items in books not yet on the statement + items on the statement not yet in books
    expected = book - outstanding + unmatched_total
    return {'book': book, 'outstanding': outstanding, 'outstanding_lines': outstanding_lines, 'unmatched_bank': unmatched_bank,
            'unmatched_total': unmatched_total, 'expected_statement': expected, 'statement': statement_balance,
            'difference': q2(statement_balance) - expected, 'can_finalise': q2(statement_balance) == expected and not unmatched_bank.exists()}


def finalise(bank: BankAccount, actor, *, as_of: date, statement_balance: Decimal, note: str = '') -> BankReconciliation:
    permissions.require(actor, 'banking', 'approve')
    with transaction.atomic():
        pos = reconciliation_position(bank, as_of, statement_balance)
        if pos['unmatched_bank'].exists():
            raise BankError('Match or ignore every bank line up to this date first.')
        if pos['difference']:
            raise BankError(f'The statement balance differs from the books by ₹{pos["difference"]:,.2f}.')
        rec = BankReconciliation.objects.create(bank_account=bank, statement_date=as_of, statement_balance=q2(statement_balance),
                                                book_balance=pos['book'], outstanding=pos['outstanding'], difference=ZERO,
                                                finalised_by=actor, note=note[:300])
        BankTransaction.objects.filter(bank_account=bank, txn_date__lte=as_of, status='matched').update(status='reconciled', reconciliation=rec)
        audit.log('bank.reconciled', rec, actor=actor, changes={'account': str(bank), 'date': as_of, 'statement': str(statement_balance),
                                                               'book': str(pos['book']), 'outstanding': str(pos['outstanding'])})
    return rec


def unreconcile(rec: BankReconciliation, actor, reason: str):
    if not permissions.has_perm(actor, 'accounting', 'approve'):
        permissions.denied(actor, 'accounting', 'approve', detail='unreconcile')
    if len(reason.strip()) < 10:
        raise BankError('Give a reason of at least 10 characters.')
    if BankReconciliation.objects.filter(bank_account=rec.bank_account, statement_date__gt=rec.statement_date).exists():
        raise BankError('Un-reconcile the later statements first.')
    from crm.core.posting import period_status
    if period_status(rec.statement_date) == 'closed':
        raise BankError('The period is closed. Reopen it first.')
    with transaction.atomic():
        BankTransaction.objects.filter(reconciliation=rec).update(status='matched', reconciliation=None)
        audit.log('bank.unreconciled', rec, actor=actor, reason=reason, changes={'date': rec.statement_date})
        rec.delete()
